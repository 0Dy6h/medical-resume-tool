#!/usr/bin/env python3
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
import re

def md_to_docx(md_path, docx_path):
    doc = Document()

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        if not line:
            i += 1
            continue

        # H1
        if line.startswith('# '):
            p = doc.add_heading(line[2:], level=1)
            p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        # H2
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=2)
        # H3
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=3)
        # Unordered list
        elif line.startswith('- ') or line.startswith('* '):
            text = line[2:]
            p = doc.add_paragraph(text, style='List Bullet')
        # Ordered list
        elif re.match(r'^\d+\.\s', line):
            text = re.sub(r'^\d+\.\s', '', line)
            p = doc.add_paragraph(text, style='List Number')
        # Code block
        elif line.startswith('```'):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                code_lines.append(lines[i].rstrip())
                i += 1
            p = doc.add_paragraph('\n'.join(code_lines))
            p.style.font.name = 'Courier New'
            p.style.font.size = Pt(9)
        # Horizontal rule
        elif line.startswith('---'):
            doc.add_paragraph('_' * 50)
        # Normal paragraph
        else:
            # Handle bold **text**
            text = line
            if '**' in text:
                p = doc.add_paragraph()
                parts = re.split(r'(\*\*.*?\*\*)', text)
                for part in parts:
                    if part.startswith('**') and part.endswith('**'):
                        run = p.add_run(part[2:-2])
                        run.bold = True
                    else:
                        p.add_run(part)
            else:
                doc.add_paragraph(text)

        i += 1

    doc.save(docx_path)
    print(f'Generated: {docx_path}')

if __name__ == '__main__':
    md_to_docx(
        'docs/beta-announcement.md',
        'docs/beta-announcement.docx'
    )
