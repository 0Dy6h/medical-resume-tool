from io import BytesIO
from pathlib import Path

import fitz
import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.main import create_app
from app.services.database import init_db, reset_db
from app.services.exporter import _find_cjk_font
from app.services.profile_import import extract_docx_text, parse_profile_from_lines


def make_client(tmp_path: Path) -> TestClient:
    db_path = tmp_path / "import-test.db"
    app = create_app(f"sqlite:///{db_path}")
    reset_db(app.state.engine)
    init_db(app.state.engine)
    return TestClient(app)


def auth_headers(client: TestClient) -> dict[str, str]:
    response = client.post("/api/auth/register", json={"username": "tester", "password": "secret123"})
    assert response.status_code == 201, response.text
    return {"Authorization": f"Bearer {response.json()['token']}"}


def sample_resume_docx() -> bytes:
    document = Document()
    document.add_heading("教育经历", level=1)
    document.add_paragraph("2021.09-2024.06 复旦大学 临床医学 硕士")
    document.add_paragraph("循证医学训练")
    document.add_paragraph("")
    document.add_paragraph("2017.09-2021.06 南京医科大学 临床医学 本科")
    document.add_heading("工作经历", level=1)
    document.add_paragraph("2023.01-2024.06 上海某三甲医院 科研助理")
    document.add_paragraph("维护随访数据库，协助统计分析")
    document.add_heading("科研项目经历", level=1)
    document.add_paragraph("慢病队列随访项目 项目成员")
    document.add_paragraph("完成 300 例随访记录核查")
    document.add_heading("发表论文", level=1)
    document.add_paragraph("慢病管理数据质量研究《中华公共卫生杂志》2024 doi:10.1234/abcd.5678")
    document.add_heading("获奖荣誉", level=1)
    document.add_paragraph("2023 优秀研究生 复旦大学 校级")
    document.add_heading("语言能力", level=1)
    document.add_paragraph("英语 CET-6")
    document.add_heading("专业技能", level=1)
    document.add_paragraph("SPSS、Python、数据管理")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def sample_resume_text() -> str:
    return "\n".join(
        [
            "教育经历",
            "2021.09-2024.06 复旦大学 临床医学 硕士",
            "循证医学训练",
            "",
            "专业技能",
            "SPSS、Python、数据管理",
        ]
    )


def sample_unsectioned_resume_text() -> str:
    return "\n".join(
        [
            "2021.09-2024.06 复旦大学 临床医学 硕士",
            "2023.01-2024.06 上海某三甲医院 临床研究中心 科研助理",
            "参与伦理材料整理、维护随访数据库、使用SPSS完成统计分析",
            "慢病队列随访项目 项目成员 完成300例随访记录核查,输出数据质量报告",
            "已取得GCP证书、大学英语六级",
            "熟悉 Python、R语言、SPSS、数据清洗",
        ]
    )


def sample_resume_pdf() -> bytes:
    font_path = _find_cjk_font()
    if font_path is None:
        pytest.skip("No CJK font available for PDF import fixture")
    document = fitz.open()
    page = document.new_page()
    page.insert_text(
        (72, 72),
        sample_resume_text(),
        fontsize=12,
        fontfile=str(font_path),
        fontname="cjk",
    )
    payload = document.tobytes()
    document.close()
    return payload


def test_parse_profile_extracts_sections_and_fields():
    lines = extract_docx_text(sample_resume_docx())
    parsed = parse_profile_from_lines(lines)

    assert len(parsed["education"]) == 2
    first_education = parsed["education"][0]
    assert first_education["school"] == "复旦大学"
    assert first_education["degree"] == "硕士"
    assert first_education["major"] == "临床医学"
    assert first_education["start"] == "2021.09"
    assert first_education["end"] == "2024.06"
    assert first_education["highlights"] == ["循证医学训练"]

    assert len(parsed["experiences"]) == 1
    experience = parsed["experiences"][0]
    assert experience["organization"] == "上海某三甲医院"
    assert experience["role"] == "科研助理"
    assert experience["highlights"]

    assert len(parsed["projects"]) == 1
    project = parsed["projects"][0]
    assert project["name"] == "慢病队列随访项目"
    assert project["role"] == "项目成员"

    assert len(parsed["publications"]) == 1
    publication = parsed["publications"][0]
    assert publication["journal"] == "中华公共卫生杂志"
    assert publication["year"] == "2024"
    assert "10.1234/abcd.5678" in publication["doi"]

    assert len(parsed["awards"]) == 1
    award = parsed["awards"][0]
    assert award["year"] == "2023"
    assert award["level"] == "校级"
    assert award["issuer"] == "复旦大学"

    assert parsed["languages"][0]["name"] == "英语"
    assert parsed["languages"][0]["level"].upper().replace("-", "") == "CET6"

    skill_names = {item["name"] for item in parsed["skills"]}
    assert {"SPSS", "Python", "数据管理"} <= skill_names


def test_parse_profile_warns_on_unrecognized_document():
    parsed = parse_profile_from_lines(["这是一段没有任何板块标题的自由文本", "第二行内容"])
    assert all(not parsed[collection] for collection in ["education", "experiences", "projects"])
    assert parsed["warnings"]


def test_import_endpoint_returns_parsed_profile(tmp_path):
    client = make_client(tmp_path)
    auth = auth_headers(client)
    response = client.post(
        "/api/profile/import",
        files={"file": ("resume.docx", sample_resume_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        headers=auth,
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["education"]
    assert payload["projects"]
    assert "warnings" in payload


def test_import_endpoint_accepts_plain_text(tmp_path):
    client = make_client(tmp_path)
    auth = auth_headers(client)
    response = client.post(
        "/api/profile/import",
        files={"file": ("resume.txt", sample_resume_text().encode("utf-8"), "text/plain")},
        headers=auth,
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["education"][0]["school"] == "复旦大学"
    assert {item["name"] for item in payload["skills"]} >= {"SPSS", "Python", "数据管理"}


def test_import_endpoint_accepts_unsectioned_plain_text(tmp_path):
    client = make_client(tmp_path)
    auth = auth_headers(client)
    response = client.post(
        "/api/profile/import",
        files={"file": ("resume.txt", sample_unsectioned_resume_text().encode("utf-8"), "text/plain")},
        headers=auth,
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["import_meta"]["extractor_name"] == "fact-extractor-v1"
    assert payload["education"][0]["school"] == "复旦大学"
    assert payload["experiences"][0]["organization"] == "上海某三甲医院"
    assert payload["projects"][0]["name"] == "慢病队列随访项目"
    assert payload["certificates"][0]["name"] == "GCP证书"
    assert payload["languages"][0]["level"] == "六级"
    assert {item["name"] for item in payload["skills"]} >= {"Python", "R语言", "SPSS", "数据清洗", "伦理", "随访"}
    assert payload["review_items"] == []
    assert payload["unassigned_blocks"] == []


def test_import_endpoint_accepts_markdown(tmp_path):
    client = make_client(tmp_path)
    auth = auth_headers(client)
    markdown = "# 教育经历\n2021.09-2024.06 复旦大学 临床医学 硕士\n\n## 专业技能\nSPSS、Python"
    response = client.post(
        "/api/profile/import",
        files={"file": ("resume.md", markdown.encode("utf-8"), "text/markdown")},
        headers=auth,
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["education"][0]["major"] == "临床医学"
    assert {item["name"] for item in payload["skills"]} >= {"SPSS", "Python"}


def test_import_endpoint_accepts_text_pdf(tmp_path):
    client = make_client(tmp_path)
    auth = auth_headers(client)
    response = client.post(
        "/api/profile/import",
        files={"file": ("resume.pdf", sample_resume_pdf(), "application/pdf")},
        headers=auth,
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["education"][0]["school"] == "复旦大学"


def test_import_endpoint_accepts_image_ocr_text(tmp_path, monkeypatch):
    client = make_client(tmp_path)
    auth = auth_headers(client)

    def fake_ocr_image_bytes(content: bytes, *, source_label: str):
        assert source_label == "resume.png"
        return sample_resume_text().splitlines(), ["OCR 低置信样例"]

    monkeypatch.setattr("app.services.profile_import.legacy._ocr_image_bytes", fake_ocr_image_bytes)
    response = client.post(
        "/api/profile/import",
        files={"file": ("resume.png", b"fake image bytes", "image/png")},
        headers=auth,
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["education"][0]["school"] == "复旦大学"
    assert "OCR 低置信样例" in payload["warnings"]


def test_import_endpoint_requires_auth(tmp_path):
    client = make_client(tmp_path)
    response = client.post(
        "/api/profile/import",
        files={"file": ("resume.docx", sample_resume_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 401


def test_import_endpoint_rejects_unsupported_file_type(tmp_path):
    client = make_client(tmp_path)
    auth = auth_headers(client)
    response = client.post("/api/profile/import", files={"file": ("resume.rtf", b"{\\rtf1}", "application/rtf")}, headers=auth)
    assert response.status_code == 400
    assert "支持" in response.text


def test_import_endpoint_rejects_corrupt_docx(tmp_path):
    client = make_client(tmp_path)
    auth = auth_headers(client)
    response = client.post(
        "/api/profile/import",
        files={"file": ("resume.docx", b"not a real docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        headers=auth,
    )
    assert response.status_code == 400
