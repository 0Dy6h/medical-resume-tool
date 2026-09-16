"""Adversarial regressions using synthetic documents and no external network."""
import asyncio
import gzip
import socket
import threading
from io import BytesIO
from types import SimpleNamespace
from zipfile import ZIP_DEFLATED, ZipFile

import httpx
import httpcore
import pytest
from fastapi.testclient import TestClient

import app.main as main
from app.services import http_client
from app.services.profile_import import legacy


def authenticated_client(tmp_path):
    client = TestClient(main.create_app(f"sqlite:///{tmp_path / 'review.db'}"), raise_server_exceptions=False)
    response = client.post("/api/auth/register", json={"username": "reviewer", "password": "Synthetic-Only-123"})
    assert response.status_code == 201
    return client, {"Authorization": "Bearer " + response.json()["token"]}


@pytest.mark.parametrize("sections", [
    [{"id": "skills", "title": "技能", "items": None}],
    [{"id": "skills", "title": "技能", "items": [None]}],
    [{"id": "skills", "title": "技能", "items": ["not an item"]}],
    [{"id": "skills", "title": "技能", "items": [{"text": {"bad": "shape"}}]}],
    [{"id": "skills", "title": "技能", "items": [{"text": "invalid\u0000text"}]}],
])
def test_malformed_draft_sections_are_rejected_before_repository(tmp_path, sections):
    client, auth = authenticated_client(tmp_path)
    response = client.put("/api/resume-drafts/99999", headers=auth, json={"sections": sections})
    assert response.status_code == 422, response.text


@pytest.mark.parametrize("filters", [{"trust": []}, {"region": []}, {"keyword": {}}, {"institution_id": "9" * 100}])
def test_report_filter_types_never_reach_sql(tmp_path, filters):
    client, auth = authenticated_client(tmp_path)
    response = client.post("/api/reports", headers=auth, json={"title": "合成边界检查", "filters": filters})
    assert response.status_code == 422, response.text


@pytest.mark.parametrize("route,body", [("/api/crawl-runs", {"institution_ids": [1]}), ("/api/institutions/1/recrawl", None)])
def test_failed_crawl_creation_releases_execution_slot(tmp_path, monkeypatch, route, body):
    from app.services.crawler import release_crawl_slot, try_acquire_crawl_slot
    client, auth = authenticated_client(tmp_path)
    def fail(*args, **kwargs):
        raise OSError("synthetic write failure")
    monkeypatch.setattr(main, "create_crawl_run", fail)
    response = client.post(route, headers=auth, json=body)
    assert response.status_code >= 500
    acquired = try_acquire_crawl_slot()
    # Always restore shared process state, including when exercising the old bug.
    release_crawl_slot()
    assert acquired, "a failed write retained the global crawl slot"


def test_unknown_length_stream_is_stopped_and_closed(monkeypatch):
    monkeypatch.setattr(http_client, "MAX_RESPONSE_BYTES", 8)
    class Body(httpx.AsyncByteStream):
        closed = False
        reads = 0
        async def __aiter__(self):
            for chunk in [b"first", b"second", b"never"]:
                self.reads += 1
                yield chunk
        async def aclose(self):
            self.closed = True
    body = Body()
    async def response(self, request):
        return httpx.Response(200, stream=body, request=request)
    monkeypatch.setattr(httpx.AsyncHTTPTransport, "handle_async_request", response)
    async def check():
        async with http_client.build_crawl_client() as client:
            with pytest.raises(http_client.OutboundBlockedError, match="响应过大"):
                await client.get("https://8.8.8.8/synthetic")
    asyncio.run(check())
    assert body.closed
    assert body.reads == 2


def test_shared_address_space_is_blocked():
    with pytest.raises(http_client.OutboundBlockedError):
        http_client.assert_public_url("http://100.64.0.1/")


def test_docx_expansion_is_checked_before_parser(monkeypatch):
    payload = BytesIO()
    with ZipFile(payload, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"a" * (2 * 1024 * 1024))
    calls = []
    def parser(*args, **kwargs):
        calls.append(True)
        return SimpleNamespace(paragraphs=[], tables=[])
    monkeypatch.setattr(legacy, "Document", parser)
    with pytest.raises(legacy.ProfileImportError, match="解压|压缩|复杂"):
        legacy.extract_docx_text(payload.getvalue())
    assert not calls


def test_gcp_knowledge_does_not_invent_certificate():
    from app.services.profile_import import build_profile_contract
    imported = build_profile_contract(["专业技能", "熟悉 GCP 规范，但尚未取得 GCP 证书"], []).to_dict()
    assert not imported["certificates"]
    assert not any(item["collection"] == "certificates" for item in imported["review_items"])


def test_crawl_thread_start_failure_is_visible_and_releases_slot(tmp_path, monkeypatch):
    from app.services.crawler import release_crawl_slot, try_acquire_crawl_slot
    client, auth = authenticated_client(tmp_path)
    start = threading.Thread.start
    def fail_worker(self):
        if self.name.startswith("resume-crawl-"):
            raise RuntimeError("synthetic thread failure")
        return start(self)
    monkeypatch.setattr(threading.Thread, "start", fail_worker)
    response = client.post("/api/crawl-runs", json={"institution_ids": [1]}, headers=auth)
    assert response.status_code == 503
    runs = client.get("/api/crawl-runs").json()
    assert runs[0]["status"] == "failed"
    assert runs[0]["errors"][0]["error_type"] == "worker_failure"
    acquired = try_acquire_crawl_slot()
    release_crawl_slot()
    assert acquired


def test_socket_connection_pins_public_resolution_and_rechecks_next_connection(monkeypatch):
    connected = []
    answers = iter(["93.184.216.34", "127.0.0.1"])
    monkeypatch.setattr(socket, "getaddrinfo", lambda *_: [(socket.AF_INET, socket.SOCK_STREAM, 6, "", (next(answers), 0))])
    async def connect(self, host, port, **kwargs):
        connected.append(host)
        return "synthetic socket"
    monkeypatch.setattr(httpcore.AnyIOBackend, "connect_tcp", connect)
    async def check():
        backend = http_client.PublicNetworkBackend()
        assert await backend.connect_tcp("rebind.example", 443) == "synthetic socket"
        with pytest.raises(http_client.OutboundBlockedError):
            await backend.connect_tcp("rebind.example", 443)
    asyncio.run(check())
    assert connected == ["93.184.216.34"]


def test_pinned_transport_keeps_original_host_and_tls_server_name(monkeypatch):
    writes, names = [], []
    class Stream(httpcore.AsyncMockStream):
        async def write(self, buffer, timeout=None):
            writes.append(buffer)
            await super().write(buffer, timeout)
        async def start_tls(self, ssl_context, server_hostname=None, timeout=None):
            names.append(server_hostname)
            return self
    stream = Stream([b"HTTP/1.1 200 OK\r\nContent-Length: 2\r\n\r\nok"])
    monkeypatch.setattr(socket, "getaddrinfo", lambda *_: [(socket.AF_INET, socket.SOCK_STREAM, 6, "", ("93.184.216.34", 0))])
    async def connect(self, *args, **kwargs):
        return stream
    monkeypatch.setattr(http_client.PublicNetworkBackend, "connect_tcp", connect)
    async def check():
        async with http_client.build_crawl_client() as client:
            response = await client.get("https://recruit.example/synthetic")
            assert response.text == "ok"
    asyncio.run(check())
    assert names == ["recruit.example"]
    assert b"Host: recruit.example\r\n" in b"".join(writes)


@pytest.mark.parametrize("size,accepted", [(32, True), (512, False)])
def test_compressed_response_budget_limits_decoded_bytes(monkeypatch, size, accepted):
    monkeypatch.setattr(http_client, "MAX_RESPONSE_BYTES", 64)
    class Body(httpx.AsyncByteStream):
        closed = False
        async def __aiter__(self):
            yield gzip.compress(b"x" * size)
        async def aclose(self):
            self.closed = True
    body = Body()
    response = httpx.Response(200, headers={"content-encoding": "gzip"}, stream=body)
    if accepted:
        asyncio.run(http_client.guard_response(response))
        assert response.content == b"x" * size
    else:
        with pytest.raises(http_client.OutboundBlockedError, match="解压"):
            asyncio.run(http_client.guard_response(response))
    assert body.closed


def test_xlsx_forged_dimensions_are_rejected_and_workbook_closed(monkeypatch):
    from openpyxl import Workbook
    from app.services import attachments
    from app.services.document_limits import DocumentLimitError
    workbook = Workbook()
    workbook.active.append(["岗位名称", "学历要求"])
    workbook.active.append(["合成岗位", "本科"])
    original = BytesIO()
    workbook.save(original)
    modified = BytesIO()
    with ZipFile(original) as source, ZipFile(modified, "w", ZIP_DEFLATED) as target:
        for name in source.namelist():
            value = source.read(name)
            if name == "xl/worksheets/sheet1.xml":
                value = value.replace(b'A1:B2', b'A1:XFD1048576')
            target.writestr(name, value)
    closed = []
    load_workbook = attachments.load_workbook
    def load(*args, **kwargs):
        result = load_workbook(*args, **kwargs)
        close = result.close
        def tracked_close():
            closed.append(True)
            close()
        result.close = tracked_close
        return result
    monkeypatch.setattr(attachments, "load_workbook", load)
    with pytest.raises(DocumentLimitError, match="行列"):
        attachments.parse_xlsx_table(modified.getvalue())
    assert closed == [True]


def test_image_pixel_budget_is_checked_before_ocr(monkeypatch):
    from PIL import Image
    monkeypatch.setattr(legacy, "MAX_IMAGE_PIXELS", 4)
    def forbid(*args):
        raise AssertionError("OCR must not start for oversized pixels")
    monkeypatch.setattr(legacy, "_image_to_string", forbid)
    image = BytesIO()
    Image.new("RGB", (3, 2)).save(image, "PNG")
    with pytest.raises(legacy.ProfileImportError, match="像素"):
        legacy.extract_profile_text("synthetic.png", image.getvalue())


@pytest.mark.parametrize("text", ["熟悉 GCP 规范", "计划取得 GCP 证书", "证书尚未取得"])
def test_import_does_not_upgrade_knowledge_or_plans_to_credentials(text):
    from app.services.profile_import import build_profile_contract
    imported = build_profile_contract([text], []).to_dict()
    assert not imported["certificates"]
    assert not any(item["collection"] == "certificates" for item in imported["review_items"])


def test_positive_certificate_and_skill_import_are_preserved():
    from app.services.profile_import import build_profile_contract
    imported = build_profile_contract(["已取得 GCP 证书", "专业技能", "不会 SPSS，熟悉 Python"], []).to_dict()
    assert imported["certificates"][0]["name"] == "GCP证书"
    names = {item["name"] for item in imported["skills"]}
    assert "Python" in names
    assert "SPSS" not in names
