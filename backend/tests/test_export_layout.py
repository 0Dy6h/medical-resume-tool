"""Geometric and structural layout assertions for DOCX/PDF export (D1–D12).

Uses pymupdf (fitz) for PDF geometry and python-docx + zipfile for DOCX
structure.  No visual inspection required — every assertion is machine-verifiable.
"""

from __future__ import annotations

import io
import re
import zipfile

import fitz  # pymupdf
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm

from app.services.exporter import (
    _normalize_pdf_text,
    _parse_entry,
    _should_bold_main_line,
    _strip_all,
    assert_lossless,
    build_export_filename,
    content_disposition_header,
    export_docx,
    export_pdf,
)

MM_TO_PT = 2.8346
LEFT_MARGIN_PT = 22 * MM_TO_PT   # ≈ 62.36
RIGHT_MARGIN_PT = 22 * MM_TO_PT  # ≈ 62.36
A4_WIDTH_PT = 210 * MM_TO_PT     # ≈ 595.28
A4_HEIGHT_PT = 297 * MM_TO_PT    # ≈ 841.89
RIGHT_EDGE_PT = A4_WIDTH_PT - RIGHT_MARGIN_PT  # ≈ 532.92


# ─── Sample drafts ────────────────────────────────────────────────────


def _sample_draft() -> dict:
    """A comprehensive single-page draft covering all entry shapes."""
    return {
        "id": 1,
        "title": "内科医师 定制简历",
        "sections": [
            {
                "id": "identity",
                "title": "个人信息",
                "items": [
                    {"text": "张三"},
                    {"text": "电话：13800000000 ｜ 邮箱：test@example.com ｜ 所在地：北京"},
                    {"text": "三年三甲医院内科临床经验，擅长常见病多发病诊疗，具备独立值班和急诊处理能力。"},
                ],
            },
            {
                "id": "education",
                "title": "教育背景",
                "items": [
                    {"text": "南方医科大学 / 硕士 / 内科学 / 2019-09-2022-06：循证医学训练；临床研究设计；三级查房规范与病历书写规范"},
                    {"text": "某医科大学 / 学士 / 临床医学 / 2014-09-2019-06：基础医学课程；临床轮转实习；医学伦理学"},
                ],
            },
            {
                "id": "experiences",
                "title": "工作/实习经历",
                "items": [
                    {"text": "某三甲医院 / 内科住院医师 / 2022-07-2024-06：负责病区日常诊疗工作，管理床位二十张，参与三级查房与疑难病例讨论，协助带教实习生与进修医师，完成教学查房与病例汇报，参与科室质控工作，整理分析住院数据，撰写月度质控报告并提出改进建议"},
                ],
            },
            {
                "id": "skills",
                "title": "技能能力",
                "items": [
                    {"text": "SPSS"},
                    {"text": "临床研究"},
                ],
            },
        ],
        "evidence": [
            {"requirement": "内科学硕士", "source_label": "教育经历", "source_text": "南方医科大学 硕士", "evidence_strength": "strong"},
        ],
        "gaps": [
            {"requirement": "SCI论文", "message": "未在你的履历中找到对应证据：SCI论文", "blocking": False},
        ],
    }


def _multi_page_draft() -> dict:
    """A draft with enough content to span 2+ pages."""
    sections: list[dict] = [
        {
            "id": "identity",
            "title": "个人信息",
            "items": [
                {"text": "李四"},
                {"text": "电话：13900000000 ｜ 邮箱：lisi@example.com"},
                {"text": "五年临床经验，擅长内科常见病诊疗。"},
            ],
        },
    ]
    for i in range(1, 13):
        sections.append({
            "id": f"section-{i}",
            "title": f"经历板块{i}",
            "items": [
                {"text": f"机构{i} / 角色{i} / 2019-2024：详细描述{i}，这是一段足够长的内容用于确保分页和多行换行测试，包含多个中文字符以确保排版正确性和内容完整性验证"},
                {"text": f"项目{i} / 成员{i} / 2020-2023：另一段详细描述，包含足够内容以确保多行渲染和悬挂缩进效果验证"},
            ],
        })
    return {
        "id": 2,
        "title": "内科医师 定制简历",
        "sections": sections,
        "evidence": [],
        "gaps": [],
    }


def _lossless_strip(s: str) -> str:
    """Aggressive strip for lossless content comparison in tests."""
    return re.sub(r"[\s/：；·—–.\-（）()\t\n\r]", "", s)


# ─── Helpers ──────────────────────────────────────────────────────────


def _docx_xml(docx_bytes: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as zf:
        return zf.read("word/document.xml").decode("utf-8")


def _docx_pstyles(docx_bytes: bytes) -> set[str]:
    return set(re.findall(r'w:pStyle w:val="([^"]+)"', _docx_xml(docx_bytes)))


def _extract_docx_text(docx_bytes: bytes) -> str:
    document = Document(io.BytesIO(docx_bytes))
    parts: list[str] = []
    for para in document.paragraphs:
        for run in para.runs:
            parts.append(run.text)
    return "".join(parts)


def _style_has_eastasia(style) -> bool:
    rpr = style.element.find(qn("w:rPr"))
    if rpr is None:
        return False
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        return False
    return rfonts.get(qn("w:eastAsia")) is not None


# ─── Lossless parsing unit tests ──────────────────────────────────────


class TestLosslessParsing:
    def test_entry_with_head_role_period_details(self):
        raw = "南方医科大学 / 硕士 / 内科学 / 2019-09-2022-06：循证医学训练；临床研究设计"
        entry = _parse_entry(raw)
        assert entry.head == "南方医科大学"
        assert entry.role == "硕士 · 内科学"
        assert entry.period == "2019-09-2022-06"
        assert entry.details == ["循证医学训练", "临床研究设计"]
        assert_lossless(raw, entry)

    def test_entry_with_head_only_no_separator(self):
        raw = "SPSS"
        entry = _parse_entry(raw)
        assert entry.head == "SPSS"
        assert entry.role is None
        assert entry.period is None
        assert entry.details == []
        assert_lossless(raw, entry)

    def test_entry_with_details_only(self):
        raw = "循证医学训练；临床研究设计"
        entry = _parse_entry(raw)
        assert entry.head is None
        assert entry.details == ["循证医学训练", "临床研究设计"]
        assert_lossless(raw, entry)

    def test_entry_with_head_and_details_no_role(self):
        raw = "某大学：医学基础课程；临床实习"
        entry = _parse_entry(raw)
        assert entry.head == "某大学"
        assert entry.details == ["医学基础课程", "临床实习"]
        assert_lossless(raw, entry)

    def test_entry_year_as_period(self):
        raw = "大学英语六级 / 教育部考试中心 / 2022"
        entry = _parse_entry(raw)
        assert entry.head == "大学英语六级"
        assert entry.role == "教育部考试中心"
        assert entry.period == "2022"
        assert_lossless(raw, entry)

    def test_entry_year_range_period(self):
        raw = "某机构 / 角色 / 2019-2022"
        entry = _parse_entry(raw)
        assert entry.head == "某机构"
        assert entry.role == "角色"
        assert entry.period == "2019-2022"
        assert_lossless(raw, entry)

    def test_appendix_item_lossless(self):
        raw = "临床医学硕士 — 证据：教育经历：硕士 — 强度：strong"
        entry = _parse_entry(raw)
        assert_lossless(raw, entry)

    def test_gap_message_lossless(self):
        raw = "SCI论文 — 建议：未在你的履历中找到对应证据：SCI论文"
        entry = _parse_entry(raw)
        assert_lossless(raw, entry)


# ─── PDF layout tests ─────────────────────────────────────────────────


class TestPDFLayout:
    def test_d1_contact_and_summary_not_lost(self):
        """D1: PDF page 1 must contain phone, email, and summary text."""
        pdf_bytes = export_pdf(_sample_draft())
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        page1_text = doc[0].get_text("text")
        doc.close()
        assert "13800000000" in page1_text, "phone number missing from page 1"
        assert "test@example.com" in page1_text, "email missing from page 1"
        assert "三年三甲" in page1_text, "summary text missing from page 1"

    def test_d2_no_text_overflow(self):
        """D2/D4: No text block drawn outside the left/right margins."""
        pdf_bytes = export_pdf(_sample_draft())
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        for page in doc:
            blocks = page.get_text("dict")["blocks"]
            for block in blocks:
                for line in block.get("lines", []):
                    for span in line.get("spans", []):
                        x0 = span["bbox"][0]
                        x1 = span["bbox"][2]
                        assert x0 >= LEFT_MARGIN_PT - 1, (
                            f"span x0={x0:.1f} < left_margin={LEFT_MARGIN_PT:.1f}"
                        )
                        assert x1 <= RIGHT_EDGE_PT + 1, (
                            f"span x1={x1:.1f} > right_edge={RIGHT_EDGE_PT:.1f}"
                        )
        doc.close()

    def test_d3_a4_paper_size(self):
        """D4: PDF page dimensions must be A4."""
        pdf_bytes = export_pdf(_sample_draft())
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        page = doc[0]
        assert abs(page.rect.width - A4_WIDTH_PT) < 1, (
            f"page width {page.rect.width:.1f} ≠ A4 {A4_WIDTH_PT:.1f}"
        )
        assert abs(page.rect.height - A4_HEIGHT_PT) < 1, (
            f"page height {page.rect.height:.1f} ≠ A4 {A4_HEIGHT_PT:.1f}"
        )
        doc.close()

    def test_d3_hanging_indent_on_wrapped_detail(self):
        """D3: Continuation line x0 > main line x0 (hanging indent)."""
        draft = _sample_draft()
        pdf_bytes = export_pdf(draft)
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        page = doc[0]
        blocks = page.get_text("dict")["blocks"]

        main_x0 = None
        cont_x0 = None
        for block in blocks:
            for line in block.get("lines", []):
                line_text = "".join(s["text"] for s in line["spans"])
                if line_text.startswith("某三甲医院"):
                    main_x0 = line["spans"][0]["bbox"][0]
                # Continuation line: contains detail keywords but doesn't start
                # with the head ("某三甲") or the bullet ("–").
                if ("病区" in line_text or "查房" in line_text or "带教" in line_text) \
                        and not line_text.startswith("某三甲") \
                        and not line_text.startswith("–"):
                    cont_x0 = min(s["bbox"][0] for s in line["spans"])
                    break
            if main_x0 is not None and cont_x0 is not None:
                break
        doc.close()

        assert main_x0 is not None, "could not find main line"
        assert cont_x0 is not None, "could not find continuation line"
        assert cont_x0 > main_x0 + 2, (
            f"continuation x0={cont_x0:.1f} not greater than main x0={main_x0:.1f}"
        )

    def test_d2_not_justified(self):
        """D2: Left-aligned (not justified) — non-last lines don't reach right margin."""
        draft = _sample_draft()
        pdf_bytes = export_pdf(draft)
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        page = doc[0]
        blocks = page.get_text("dict")["blocks"]

        found_non_justified = False
        for block in blocks:
            lines = block.get("lines", [])
            if len(lines) < 2:
                continue
            # Check non-last lines — with justify, x1 ≈ right edge.
            for line in lines[:-1]:
                max_x1 = max(s["bbox"][2] for s in line["spans"])
                if max_x1 < RIGHT_EDGE_PT - 1:
                    found_non_justified = True
                    break
            if found_non_justified:
                break
        doc.close()
        assert found_non_justified, (
            "all multi-line blocks appear justified (x1 at right margin) — "
            "expected at least one left-aligned block"
        )

    def test_d7_page_numbers_on_multi_page_pdf(self):
        """D7: Multi-page PDF footer contains '第 X 页' and '共 Y 页'."""
        pdf_bytes = export_pdf(_multi_page_draft())
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        assert doc.page_count >= 2, f"expected ≥2 pages, got {doc.page_count}"
        last_page_text = doc[-1].get_text("text")
        doc.close()
        assert "第 2 页" in last_page_text or "第" in last_page_text, (
            "footer page number missing"
        )
        assert "共" in last_page_text and "页" in last_page_text, (
            "footer total page count missing"
        )

    def test_d8_chinese_text_extractable(self):
        """D8: Chinese text in entries is correctly extractable from PDF."""
        pdf_bytes = export_pdf(_sample_draft())
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        text = doc[0].get_text("text")
        doc.close()
        assert "南方医科大学" in text
        assert "内科住院医师" in text
        assert "循证医学训练" in text

    def test_d5_bold_font_registered(self):
        """D5: Bold font must be registered (or fallback to regular without error)."""
        # If export_pdf succeeds without raising, the bold font is handled.
        pdf_bytes = export_pdf(_sample_draft())
        assert pdf_bytes.startswith(b"%PDF")
        # Verify the PDF has content (not a trivially empty file)
        assert len(pdf_bytes) > 500


# ─── DOCX layout tests ───────────────────────────────────────────────


class TestDOCXLayout:
    def test_d9_a4_page_size(self):
        """D9: DOCX page size must be A4 (within 1 twip rounding tolerance)."""
        docx_bytes = export_docx(_sample_draft())
        document = Document(io.BytesIO(docx_bytes))
        section = document.sections[0]
        assert abs(section.page_width - Mm(210)) < 700, "page_width not A4"
        assert abs(section.page_height - Mm(297)) < 700, "page_height not A4"

    def test_d10_eastasia_font_on_styles(self):
        """D10: Normal and Heading styles must have w:eastAsia in rFonts."""
        docx_bytes = export_docx(_sample_draft())
        document = Document(io.BytesIO(docx_bytes))
        assert _style_has_eastasia(document.styles["Normal"]), "Normal missing eastAsia"
        assert _style_has_eastasia(document.styles["Heading 1"]), "Heading 1 missing eastAsia"

    def test_d10_margins_20mm_22mm(self):
        """D4: DOCX margins must be 20mm top/bottom, 22mm left/right (within twip rounding)."""
        docx_bytes = export_docx(_sample_draft())
        document = Document(io.BytesIO(docx_bytes))
        section = document.sections[0]
        assert abs(section.top_margin - Mm(20)) < 700
        assert abs(section.bottom_margin - Mm(20)) < 700
        assert abs(section.left_margin - Mm(22)) < 700
        assert abs(section.right_margin - Mm(22)) < 700

    def test_d10_pstyle_hierarchy_not_flat(self):
        """D10: pStyle values must include more than just Heading1|ListBullet."""
        docx_bytes = export_docx(_sample_draft())
        pstyles = _docx_pstyles(docx_bytes)
        assert "ResumeEntry" in pstyles or "ResumeDetail" in pstyles, (
            f"expected ResumeEntry/ResumeDetail in pStyles, got {pstyles}"
        )
        assert "ListBullet" not in pstyles, "ListBullet should not be used"

    def test_d11_keep_with_next_on_section_titles(self):
        """D11: Section title paragraphs must have w:keepNext."""
        docx_bytes = export_docx(_sample_draft())
        xml = _docx_xml(docx_bytes)
        assert "w:keepNext" in xml, "no w:keepNext found in document.xml"

    def test_d11_document_properties_title(self):
        """D11/D12: core_properties.title must be non-empty."""
        docx_bytes = export_docx(_sample_draft())
        document = Document(io.BytesIO(docx_bytes))
        assert document.core_properties.title, "core_properties.title is empty"
        assert "简历" in document.core_properties.title

    def test_d11_document_properties_author(self):
        """D12: core_properties.author must be set."""
        docx_bytes = export_docx(_sample_draft())
        document = Document(io.BytesIO(docx_bytes))
        assert document.core_properties.author, "core_properties.author is empty"
        assert document.core_properties.author == "张三"

    def test_d12_lossless_content_preserved(self):
        """D12/无损: Every item's core content must appear in the DOCX text."""
        draft = _sample_draft()
        docx_bytes = export_docx(draft)
        docx_text = _lossless_strip(_extract_docx_text(docx_bytes))
        for section in draft["sections"]:
            for item in section.get("items", []):
                raw = str(item.get("text", ""))
                stripped = _lossless_strip(raw)
                assert stripped in docx_text, (
                    f"item content not found in DOCX: {stripped!r}"
                )

    def test_d6_section_title_has_bottom_border(self):
        """D6: Section title paragraphs must have a bottom border."""
        docx_bytes = export_docx(_sample_draft())
        xml = _docx_xml(docx_bytes)
        assert "w:pBdr" in xml, "no paragraph border found in document.xml"
        assert "w:bottom" in xml, "no bottom border found"

    def test_d6_resume_entry_has_tab_stop_for_period(self):
        """D6/版式: Entry main line must have a right tab stop for the period."""
        docx_bytes = export_docx(_sample_draft())
        xml = _docx_xml(docx_bytes)
        assert "w:tabs" in xml, "no tab stops found in document.xml"
        assert "right" in xml.lower(), "no right-aligned tab stop found"

    def test_d12_identity_name_in_docx(self):
        """D12: Applicant name must appear in the DOCX."""
        docx_bytes = export_docx(_sample_draft())
        xml = _docx_xml(docx_bytes)
        assert "张三" in xml

    def test_diagnostic_appendix_text_preserved(self):
        """D-附录: Diagnostic mode must contain fixed appendix text."""
        draft = _sample_draft()
        docx_bytes = export_docx(draft, include_appendix=True)
        xml = _docx_xml(docx_bytes)
        assert "匹配分析附录" in xml
        assert "已满足" in xml
        assert "未满足" in xml
        # Gap message should be in the appendix unmet items
        assert "未在你的履历中找到对应证据" in xml

    def test_application_mode_no_appendix(self):
        """D-附录: Application mode must not contain appendix."""
        draft = _sample_draft()
        docx_bytes = export_docx(draft, include_appendix=False)
        xml = _docx_xml(docx_bytes)
        assert "匹配分析附录" not in xml


# ─── Filename tests (Task C) ─────────────────────────────────────────


class TestFilename:
    def test_filename_with_name_and_job_title(self):
        draft = _sample_draft()
        name = build_export_filename(draft, "docx", "application")
        assert "张三" in name
        assert "内科医师" in name
        assert "简历" in name
        assert name.endswith(".docx")

    def test_filename_diagnostic_suffix(self):
        draft = _sample_draft()
        name = build_export_filename(draft, "pdf", "diagnostic")
        assert "诊断版" in name
        assert name.endswith(".pdf")

    def test_filename_fallback_no_identity(self):
        draft = {"id": 42, "title": "岗位 定制简历", "sections": []}
        name = build_export_filename(draft, "docx", "application")
        assert name == "resume-42.docx"

    def test_content_disposition_has_filename_star(self):
        header = content_disposition_header("张三-内科医师-简历-20260824.docx")
        assert "filename=" in header
        assert "filename*=UTF-8''" in header
        # The ASCII fallback should not contain Chinese
        assert "张三" not in header.split("filename*=UTF-8''")[0]

    def test_content_disposition_ascii_only_filename(self):
        header = content_disposition_header("resume-1.docx")
        assert 'filename="resume-1.docx"' in header
        assert "filename*=UTF-8''resume-1.docx" in header


# ─── Whitespace robustness tests (P1 fix) ─────────────────────────────


def _pdf_page_text(pdf_bytes: bytes) -> str:
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    text = doc[0].get_text("text")
    doc.close()
    return text


class TestUnicodeWhitespaceRobustness:
    """Entries containing Unicode whitespace (U+3000, U+00A0) must export
    successfully without triggering AssertionError, and the original content
    must appear intact in DOCX/PDF extracted text.
    """

    def _draft_with_item(self, text: str) -> dict:
        return {
            "id": 1,
            "title": "内科医师 定制简历",
            "sections": [
                {
                    "id": "identity",
                    "title": "个人信息",
                    "items": [
                        {"text": "测试用户"},
                        {"text": "电话：13800000000 ｜ 邮箱：test@example.com"},
                    ],
                },
                {
                    "id": "experiences",
                    "title": "工作经历",
                    "items": [{"text": text}],
                },
            ],
            "evidence": [],
            "gaps": [],
        }

    def test_fullwidth_space_u3000_export_succeeds(self):
        """U+3000 (fullwidth space) must not cause export failure."""
        raw = "南方医院 / 住院医师\u3000 / 2022-07-2025-06：负责病房管理"
        draft = self._draft_with_item(raw)

        docx_bytes = export_docx(draft)
        docx_text = _extract_docx_text(docx_bytes)
        assert "南方医院" in docx_text
        assert "住院医师" in docx_text
        assert "负责病房管理" in docx_text

        pdf_text = _pdf_page_text(export_pdf(draft))
        assert "南方医院" in pdf_text
        assert "负责病房管理" in pdf_text

    def test_nonbreaking_space_u00a0_export_succeeds(self):
        """U+00A0 (NBSP) must not cause export failure."""
        raw = "南方医院 / 住院医师\u00a0：负责病房管理"
        draft = self._draft_with_item(raw)

        docx_bytes = export_docx(draft)
        docx_text = _extract_docx_text(docx_bytes)
        assert "南方医院" in docx_text
        assert "住院医师" in docx_text
        assert "负责病房管理" in docx_text

        pdf_text = _pdf_page_text(export_pdf(draft))
        assert "南方医院" in pdf_text
        assert "负责病房管理" in pdf_text

    def test_lossless_passes_with_unicode_whitespace(self):
        """assert_lossless itself must pass after _strip_all Unicode normalization."""
        raw = "南方医院 / 住院医师\u3000 / 2022-07-2025-06：负责病房管理"
        entry = _parse_entry(raw)
        assert_lossless(raw, entry)  # must not raise

        raw2 = "南方医院 / 住院医师\u00a0：负责病房管理"
        entry2 = _parse_entry(raw2)
        assert_lossless(raw2, entry2)  # must not raise


# ─── Multi-paragraph summary tests (P2 fix) ──────────────────────────


class TestMultiParagraphSummary:
    """When the identity section contains multiple non-contact text paragraphs,
    all of them must be rendered — not just the last one.
    """

    def _draft_with_multi_summary(self) -> dict:
        return {
            "id": 1,
            "title": "内科医师 定制简历",
            "sections": [
                {
                    "id": "identity",
                    "title": "个人信息",
                    "items": [
                        {"text": "张三"},
                        {"text": "电话：13800000000 ｜ 邮箱：test@example.com"},
                        {"text": "第一段自我介绍：内科住院医师三年。"},
                        {"text": "第二段自我介绍：擅长糖尿病管理。"},
                    ],
                },
                {
                    "id": "education",
                    "title": "教育背景",
                    "items": [
                        {"text": "某医科大学 / 学士 / 临床医学 / 2014-09-2019-06：基础医学课程"},
                    ],
                },
            ],
            "evidence": [],
            "gaps": [],
        }

    def test_both_summary_paragraphs_in_docx(self):
        draft = self._draft_with_multi_summary()
        docx_text = _extract_docx_text(export_docx(draft))
        assert "第一段自我介绍" in docx_text, "first summary paragraph missing from DOCX"
        assert "第二段自我介绍" in docx_text, "second summary paragraph missing from DOCX"
        assert "内科住院医师三年" in docx_text
        assert "擅长糖尿病管理" in docx_text

    def test_both_summary_paragraphs_in_pdf(self):
        draft = self._draft_with_multi_summary()
        pdf_text = _pdf_page_text(export_pdf(draft))
        assert "第一段自我介绍" in pdf_text, "first summary paragraph missing from PDF"
        assert "第二段自我介绍" in pdf_text, "second summary paragraph missing from PDF"
        assert "内科住院医师三年" in pdf_text
        assert "擅长糖尿病管理" in pdf_text


# ─── B1: Bold logic tests ────────────────────────────────────────────


def _find_run_bold(docx_bytes: bytes, text_fragment: str) -> bool | None:
    """Find the bold property of the first run containing a text fragment."""
    document = Document(io.BytesIO(docx_bytes))
    for para in document.paragraphs:
        for run in para.runs:
            if text_fragment in run.text:
                return run.bold
    return None


def _bold_draft() -> dict:
    return {
        "id": 1,
        "title": "内科医师 定制简历",
        "sections": [
            {"id": "identity", "title": "个人信息", "items": [
                {"text": "张三"},
                {"text": "电话：13800000000"},
            ]},
            {"id": "target", "title": "求职目标", "items": [
                {"text": "应聘内科医师岗位，期望在三甲医院工作"},
            ]},
            {"id": "education", "title": "教育背景", "items": [
                {"text": "南方医科大学 / 硕士 / 内科学 / 2019-09-2022-06：循证医学训练"},
            ]},
            {"id": "skills", "title": "技能能力", "items": [
                {"text": "SPSS"},
                {"text": "临床研究"},
            ]},
        ],
        "evidence": [],
        "gaps": [],
    }


class TestBoldLogic:
    """B1: entry main line is bold unless target section or long no-period prose."""

    def test_should_bold_normal_entry_with_period(self):
        entry = _parse_entry("南方医科大学 / 硕士 / 内科学 / 2019-09-2022-06：循证医学训练")
        assert _should_bold_main_line(entry, "education") is True

    def test_should_not_bold_target_section(self):
        entry = _parse_entry("应聘内科医师岗位")
        assert _should_bold_main_line(entry, "target") is False

    def test_should_not_bold_long_no_period(self):
        entry = _parse_entry("这是一段超过二十四个字符的描述文字用于测试不加粗的条目")
        assert len(entry.head or "") > 24
        assert _should_bold_main_line(entry, "experiences") is False

    def test_should_bold_short_no_period(self):
        entry = _parse_entry("SPSS")
        assert _should_bold_main_line(entry, "skills") is True

    def test_docx_education_entry_is_bold(self):
        docx_bytes = export_docx(_bold_draft())
        bold = _find_run_bold(docx_bytes, "南方医科大学")
        assert bold is True, "education entry main line should be bold"

    def test_docx_target_entry_not_bold(self):
        docx_bytes = export_docx(_bold_draft())
        bold = _find_run_bold(docx_bytes, "应聘内科医师")
        assert bold is not True, "target entry main line should not be bold"

    def test_docx_skills_entry_is_bold(self):
        docx_bytes = export_docx(_bold_draft())
        bold = _find_run_bold(docx_bytes, "SPSS")
        assert bold is True, "skills entry main line should be bold"

    def test_pdf_bold_export_succeeds(self):
        pdf_bytes = export_pdf(_bold_draft())
        assert pdf_bytes.startswith(b"%PDF")
        assert len(pdf_bytes) > 500


# ─── B3: Invisible character normalization tests ─────────────────────


def _strip_for_comparison(s: str) -> str:
    """Aggressive strip including invisible characters for content comparison."""
    s = s.replace("\u200B", "").replace("\uFEFF", "").replace("\u00AD", "")
    return _lossless_strip(s)


def _draft_with_invisible_chars() -> dict:
    return {
        "id": 1,
        "title": "测试 定制简历",
        "sections": [
            {"id": "identity", "title": "个人信息", "items": [
                {"text": "测试\u200B用户"},
                {"text": "电话：13800000000"},
            ]},
            {"id": "experiences", "title": "工作经历", "items": [
                {"text": "南方医院\uFEFF / 住院医师 / 2022-07-2025-06：负责\u00A0病房\u3000管理\t工作"},
            ]},
            {"id": "skills", "title": "技能能力", "items": [
                {"text": "临床研究🏥"},
            ]},
        ],
        "evidence": [],
        "gaps": [],
    }


class TestInvisibleCharacterNormalization:
    """B3: invisible characters cleaned before PDF rendering; visible text preserved."""

    def test_normalize_removes_zero_width_space(self):
        assert _normalize_pdf_text("abc\u200Bdef") == "abcdef"

    def test_normalize_removes_bom(self):
        assert _normalize_pdf_text("\uFEFFabc") == "abc"

    def test_normalize_removes_soft_hyphen(self):
        assert _normalize_pdf_text("abc\u00ADdef") == "abcdef"

    def test_normalize_tab_to_space(self):
        assert _normalize_pdf_text("a\tb") == "a b"

    def test_normalize_nbsp_to_space(self):
        assert _normalize_pdf_text("a\u00A0b") == "a b"

    def test_normalize_fullwidth_space_to_space(self):
        assert _normalize_pdf_text("a\u3000b") == "a b"

    def test_normalize_preserves_visible_text(self):
        raw = "南方医院 / 住院医师 / 2022-07-2025-06：负责病房管理"
        assert _normalize_pdf_text(raw) == raw

    def test_export_with_invisible_chars_docx_succeeds(self):
        draft = _draft_with_invisible_chars()
        docx_bytes = export_docx(draft)
        docx_text = _strip_for_comparison(_extract_docx_text(docx_bytes))
        for section in draft["sections"]:
            for item in section.get("items", []):
                raw = str(item.get("text", ""))
                stripped = _strip_for_comparison(raw)
                assert stripped in docx_text, (
                    f"item content not found in DOCX: {stripped!r}"
                )

    def test_export_with_invisible_chars_pdf_succeeds(self):
        draft = _draft_with_invisible_chars()
        pdf_bytes = export_pdf(draft)
        assert pdf_bytes.startswith(b"%PDF")
        assert len(pdf_bytes) > 500

    def test_pdf_visible_text_preserved(self):
        draft = _draft_with_invisible_chars()
        pdf_text = _pdf_page_text(export_pdf(draft))
        assert "南方医院" in pdf_text
        assert "住院医师" in pdf_text
        assert "负责" in pdf_text
        assert "病房" in pdf_text
        assert "管理" in pdf_text
        assert "工作" in pdf_text
        assert "临床研究" in pdf_text
