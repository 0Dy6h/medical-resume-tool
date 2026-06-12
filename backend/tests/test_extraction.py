"""测试增强的文档提取功能。"""
from io import BytesIO

import pytest
from docx import Document

from app.services.profile_import.extraction import (
    extract_docx_document,
    extract_markdown_document,
)


def test_extract_docx_document_with_paragraphs():
    doc = Document()
    doc.add_paragraph("复旦大学 临床医学 硕士 2021.09-2024.06")
    doc.add_paragraph("上海某三甲医院 科研助理 2023.01-2024.06")

    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)

    result = extract_docx_document(bio.read())
    assert len(result.blocks) == 2
    assert result.blocks[0].text == "复旦大学 临床医学 硕士 2021.09-2024.06"
    assert result.blocks[0].kind == "paragraph"
    assert result.text_quality == 0.9


def test_extract_docx_document_with_table():
    doc = Document()
    table = doc.add_table(rows=2, cols=3)
    table.rows[0].cells[0].text = "复旦大学"
    table.rows[0].cells[1].text = "临床医学"
    table.rows[0].cells[2].text = "硕士"

    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)

    result = extract_docx_document(bio.read())
    assert len(result.blocks) >= 1
    # 表格行被合并成一行
    assert any("复旦大学" in block.text for block in result.blocks)
    assert any(block.kind == "table_row" for block in result.blocks)


def test_extract_markdown_document_with_headings():
    content = """# 教育经历
复旦大学 临床医学 硕士

## 工作经历
上海某三甲医院 科研助理
""".encode("utf-8")

    result = extract_markdown_document(content)
    headings = [b for b in result.blocks if b.kind == "heading"]
    assert len(headings) == 2
    assert headings[0].text == "教育经历"
    assert headings[1].text == "工作经历"


def test_extract_markdown_document_with_bullets():
    content = """- 参与伦理材料整理
- 维护随访数据库
- 使用 SPSS 完成统计分析
""".encode("utf-8")

    result = extract_markdown_document(content)
    bullets = [b for b in result.blocks if b.kind == "bullet"]
    assert len(bullets) == 3
    assert bullets[0].text == "参与伦理材料整理"
    assert bullets[1].text == "维护随访数据库"


def test_extract_markdown_handles_mixed_content():
    content = """# 教育经历
2021.09-2024.06 复旦大学 临床医学 硕士

## 项目经历
慢病队列随访项目
- 完成300例随访记录核查
- 输出数据质量报告
""".encode("utf-8")

    result = extract_markdown_document(content)
    assert len(result.blocks) >= 5
    kinds = {b.kind for b in result.blocks}
    assert "heading" in kinds
    assert "bullet" in kinds
    assert "paragraph" in kinds
