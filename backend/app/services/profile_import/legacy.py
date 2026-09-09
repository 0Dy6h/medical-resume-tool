"""履历资料导入：文档/文本/图片提取 + 规则解析。

解析是无状态的：上传文件只在内存中转成行文本，再按板块标题切分、逐条目抽取结构化字段。
识别不出的内容保留在 highlights 或 warnings 中，不丢弃。图片 OCR 依赖本机 Tesseract；不可用时
返回 warning，而不是让整个导入失败。
"""
from __future__ import annotations

import os
import re
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document

from app.services.classifier import normalize_text


SECTION_ALIASES: dict[str, list[str]] = {
    "education": ["教育经历", "教育背景", "学习经历", "学历信息"],
    "experiences": ["工作经历", "实习经历", "工作/实习经历", "工作与实习经历", "职业经历", "工作实习经历"],
    "projects": ["科研项目", "项目经历", "科研经历", "课题经历", "科研/项目经历", "科研项目经历", "课题"],
    "publications": ["发表论文", "论文发表", "学术成果", "学术论文", "论文", "出版物"],
    "certificates": ["证书资质", "资格证书", "资质证书", "证书"],
    "skills": ["专业技能", "技能能力", "技能特长", "技能"],
    "teaching": ["教学经历", "授课经历", "教学工作"],
    "awards": ["获奖荣誉", "荣誉奖项", "获奖经历", "获奖情况", "奖励荣誉", "获奖", "荣誉"],
    "languages": ["语言能力", "语言水平", "外语能力", "外语水平", "语言"],
}

_ALIAS_TO_COLLECTION = {
    alias: collection for collection, aliases in SECTION_ALIASES.items() for alias in aliases
}

_HEADING_DECORATION = re.compile(r"^[\s\d一二三四五六七八九十、.．()（）\[\]【】*#:：\-—]+|[\s:：、.．*#\-—]+$")
_BULLET_PREFIX = re.compile(r"^[•·●◦▪‣*\-–—]+\s*|^\d{1,2}[.、)）]\s*")
_DATE_RANGE = re.compile(
    r"(?P<start>\d{4}(?:[年./-]\d{1,2})?月?)\s*[-–—~～至到]+\s*(?P<end>至今|现在|\d{4}(?:[年./-]\d{1,2})?月?)"
)
_YEAR = re.compile(r"(?:19|20)\d{2}")
_YEAR_LEAD = re.compile(r"^\d{4}")
_DOI = re.compile(r"10\.\d{4,9}/\S+")
_JOURNAL = re.compile(r"《([^》]+)》")
_DEGREE = re.compile(r"博士|硕士|学士|本科|大专")

_SCHOOL_SUFFIXES = ["大学", "学院", "学校", "医学部"]
_ORG_SUFFIXES = ["医院", "大学", "学院", "中心", "研究所", "公司", "集团", "卫生院", "疾控", "诊所", "实验室", "科室"]
_ROLE_SUFFIXES = [
    "助理", "医师", "护师", "护士", "研究员", "工程师", "专员", "主管", "经理",
    "组长", "干事", "秘书", "实习生", "技师", "药师", "讲师", "教师", "助教", "成员", "负责人",
]
_PROJECT_ROLE_SUFFIXES = ["负责人", "主持", "成员", "参与", "骨干"]
_PROJECT_NAME_HINTS = ["项目", "课题", "研究", "队列", "平台", "基金"]
_ISSUER_SUFFIXES = ["部", "局", "委员会", "协会", "学会", "中心", "大学", "学院", "医院", "厅", "办公室"]
_AWARD_LEVEL = re.compile(r"国家级|省部级|省级|市级|校级|院级|特等奖|一等奖|二等奖|三等奖|金奖|银奖|铜奖|优秀奖")
_LANGUAGE_NAME = re.compile(r"英语|日语|法语|德语|俄语|韩语|西班牙语|阿拉伯语|普通话")
_LANGUAGE_LEVEL = re.compile(r"CET-?[46]|四级|六级|专业四级|专业八级|专四|专八|N[1-5]|雅思\s*[\d.]*|托福\s*\d*|IELTS\s*[\d.]*|TOEFL\s*\d*")
_TOKEN_SPLIT = re.compile(r"[\s,，;；、|/]+")

TEXT_EXTENSIONS = {".txt", ".md", ".markdown"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
SUPPORTED_IMPORT_EXTENSIONS = {".docx", ".pdf", *TEXT_EXTENSIONS, *IMAGE_EXTENSIONS}
MAX_PDF_PAGES = int(os.getenv("PROFILE_IMPORT_MAX_PDF_PAGES", "20"))
MAX_PDF_OCR_PAGES = int(os.getenv("PROFILE_IMPORT_MAX_PDF_OCR_PAGES", "3"))
MAX_IMAGE_FRAMES = int(os.getenv("PROFILE_IMPORT_MAX_IMAGE_FRAMES", "5"))
OCR_LANGUAGES = os.getenv("PROFILE_IMPORT_OCR_LANGUAGES", "chi_sim+eng")


@dataclass
class ProfileTextExtraction:
    lines: list[str]
    warnings: list[str] = field(default_factory=list)


class ProfileImportError(ValueError):
    pass


def extract_profile_text(filename: str, content: bytes) -> ProfileTextExtraction:
    extension = Path(filename or "").suffix.lower()
    if extension not in SUPPORTED_IMPORT_EXTENSIONS:
        supported = "、".join(sorted(SUPPORTED_IMPORT_EXTENSIONS))
        raise ProfileImportError(f"仅支持以下资料格式：{supported}")
    if not content:
        raise ProfileImportError("上传文件为空")
    if extension == ".docx":
        return ProfileTextExtraction(extract_docx_text(content))
    if extension == ".pdf":
        return extract_pdf_text(content)
    if extension in TEXT_EXTENSIONS:
        # markdown 标题标记在行式管线里没有语义，还会让「# 姓名」因前缀
        # 不满足纯中文姓名检测而被静默丢弃（docx 标题段落无此前缀，行为一致）。
        return extract_plain_text(content, strip_markdown_headings=extension != ".txt")
    return extract_image_text(content, filename or extension)


def extract_docx_text(content: bytes) -> list[str]:
    try:
        document = Document(BytesIO(content))
    except Exception as exc:
        raise ProfileImportError("无法读取该 docx 文件，请确认文件未损坏") from exc
    lines: list[str] = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [normalize_text(cell.text) for cell in row.cells]
            unique = [cell for cell in dict.fromkeys(cells) if cell]
            if unique:
                lines.append("；".join(unique))
    return lines


def extract_pdf_text(content: bytes) -> ProfileTextExtraction:
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:  # pragma: no cover - dependency is declared in pyproject.
        raise ProfileImportError("PDF 导入需要安装 PyMuPDF") from exc

    try:
        document = fitz.open(stream=content, filetype="pdf")
    except Exception as exc:
        raise ProfileImportError("无法读取该 PDF 文件，请确认文件未损坏") from exc

    lines: list[str] = []
    warnings: list[str] = []
    textless_pages: list[int] = []
    try:
        page_count = document.page_count
        limit = min(page_count, MAX_PDF_PAGES)
        if page_count > MAX_PDF_PAGES:
            warnings.append(f"PDF 共 {page_count} 页，仅解析前 {MAX_PDF_PAGES} 页")
        for page_index in range(limit):
            page = document.load_page(page_index)
            page_text = page.get_text("text")
            if normalize_text(page_text):
                lines.extend(page_text.splitlines())
                lines.append("")
            else:
                textless_pages.append(page_index)
        if textless_pages:
            ocr_lines, ocr_warnings = _ocr_pdf_pages(document, textless_pages)
            if ocr_lines:
                lines.extend(ocr_lines)
            warnings.extend(ocr_warnings)
        if not any(normalize_text(line) for line in lines):
            warnings.append("PDF 中未提取到可识别文本；如果是扫描件，请确认本机已安装 Tesseract OCR")
    finally:
        document.close()
    return ProfileTextExtraction(lines=lines, warnings=_dedupe(warnings))


_MARKDOWN_HEADING = re.compile(r"^#{1,6}\s+")


def extract_plain_text(content: bytes, *, strip_markdown_headings: bool = False) -> ProfileTextExtraction:
    warnings: list[str] = []

    def _decode() -> tuple[list[str], bool]:
        for encoding in ["utf-8-sig", "utf-8", "gb18030", "gbk", "big5"]:
            try:
                return content.decode(encoding).splitlines(), True
            except UnicodeDecodeError:
                continue
        return content.decode("utf-8", errors="replace").splitlines(), False

    lines, decoded = _decode()
    if strip_markdown_headings:
        lines = [_MARKDOWN_HEADING.sub("", line) for line in lines]
    if not decoded:
        warnings.append("文本编码无法完全识别，已尽量保留可读内容")
    return ProfileTextExtraction(lines, warnings)


def extract_image_text(content: bytes, filename: str) -> ProfileTextExtraction:
    lines, warnings = _ocr_image_bytes(content, source_label=filename)
    if not any(normalize_text(line) for line in lines):
        warnings.append("图片中未识别到可解析文本；请确认图片清晰，或改用 docx/pdf/txt 上传")
    return ProfileTextExtraction(lines=lines, warnings=_dedupe(warnings))


def _ocr_pdf_pages(document: Any, page_indexes: list[int]) -> tuple[list[str], list[str]]:
    if not page_indexes:
        return [], []
    ocr_indexes = page_indexes[:MAX_PDF_OCR_PAGES]
    warnings = []
    if len(page_indexes) > MAX_PDF_OCR_PAGES:
        warnings.append(f"PDF 有 {len(page_indexes)} 页未含可提取文本，仅尝试 OCR 前 {MAX_PDF_OCR_PAGES} 页")
    lines: list[str] = []
    for page_index in ocr_indexes:
        page = document.load_page(page_index)
        pixmap = page.get_pixmap(matrix=_pdf_ocr_matrix(), alpha=False)
        page_lines, page_warnings = _ocr_image_bytes(pixmap.tobytes("png"), source_label=f"PDF 第 {page_index + 1} 页")
        lines.extend(page_lines)
        warnings.extend(page_warnings)
    return lines, _dedupe(warnings)


def _pdf_ocr_matrix() -> Any:
    import fitz

    return fitz.Matrix(2, 2)


def _ocr_image_bytes(content: bytes, *, source_label: str) -> tuple[list[str], list[str]]:
    try:
        from PIL import Image, ImageOps, ImageSequence
        import pytesseract
    except ImportError as exc:  # pragma: no cover - dependencies are declared in pyproject.
        return [], [f"{source_label} 需要图片 OCR 依赖：{exc}"]

    tesseract_cmd = os.getenv("TESSERACT_CMD")
    if tesseract_cmd:
        pytesseract.pytesseract.tesseract_cmd = tesseract_cmd
    try:
        image = Image.open(BytesIO(content))
    except Exception as exc:
        raise ProfileImportError("无法读取该图片文件，请确认文件未损坏") from exc

    lines: list[str] = []
    warnings: list[str] = []
    try:
        frames = list(ImageSequence.Iterator(image))[:MAX_IMAGE_FRAMES]
        if getattr(image, "n_frames", 1) > MAX_IMAGE_FRAMES:
            warnings.append(f"{source_label} 包含多帧图片，仅 OCR 前 {MAX_IMAGE_FRAMES} 帧")
        for index, frame in enumerate(frames, start=1):
            try:
                prepared = ImageOps.exif_transpose(frame).convert("RGB")
            except Exception as exc:
                warnings.append(f"{source_label} 图片帧无法解码：{exc}")
                break
            try:
                text = _image_to_string(pytesseract, prepared)
            except Exception as exc:
                warnings.append(f"{source_label} OCR 不可用：{_friendly_ocr_error(exc)}")
                break
            if normalize_text(text):
                lines.extend(text.splitlines())
                if len(frames) > 1 and index < len(frames):
                    lines.append("")
    except Exception as exc:
        raise ProfileImportError("无法读取该图片文件，请确认文件未损坏") from exc
    finally:
        image.close()
    return lines, _dedupe(warnings)


def _image_to_string(pytesseract_module: Any, image: Any) -> str:
    try:
        return pytesseract_module.image_to_string(image, lang=OCR_LANGUAGES)
    except Exception:
        if OCR_LANGUAGES != "eng":
            return pytesseract_module.image_to_string(image, lang="eng")
        raise


def _friendly_ocr_error(exc: Exception) -> str:
    message = str(exc).strip()
    if "tesseract is not installed" in message.lower() or "not in your path" in message.lower():
        return "未找到 Tesseract 可执行程序，请安装 Tesseract 或配置 TESSERACT_CMD"
    if "failed loading language" in message.lower() or "couldn't load any languages" in message.lower():
        return f"未安装 OCR 语言包 {OCR_LANGUAGES}，请安装中文/英文语言包或调整 PROFILE_IMPORT_OCR_LANGUAGES"
    return message or exc.__class__.__name__


def _dedupe(items: list[str]) -> list[str]:
    return list(dict.fromkeys(item for item in items if item))


def parse_profile_from_lines(lines: list[str]) -> dict[str, Any]:
    result: dict[str, Any] = {collection: [] for collection in SECTION_ALIASES}
    warnings: list[str] = []
    unassigned: list[str] = []
    current: str | None = None
    buffer: list[str] = []

    for raw in lines:
        line = normalize_text(raw)
        collection = _match_heading(line)
        if collection:
            if current:
                result[current].extend(_build_entries(current, buffer))
            current = collection
            buffer = []
        elif current:
            buffer.append(line)
        elif line:
            unassigned.append(line)
    if current:
        result[current].extend(_build_entries(current, buffer))

    if unassigned:
        preview = "；".join(unassigned[:3])
        warnings.append(f"有 {len(unassigned)} 行内容未能归入任何板块：{preview}")
    if not any(result[collection] for collection in SECTION_ALIASES):
        warnings.append("未从文档中识别出履历条目，请检查是否包含「教育经历」「项目经历」等板块标题")
    result["warnings"] = warnings
    return result


def _match_heading(line: str) -> str | None:
    if not line or len(line) > 12:
        return None
    cleaned = _HEADING_DECORATION.sub("", line)
    return _ALIAS_TO_COLLECTION.get(cleaned)


def _build_entries(collection: str, lines: list[str]) -> list[dict[str, Any]]:
    if collection in {"education", "experiences", "projects"}:
        blocks = _split_blocks(lines)
        builder = {
            "education": _build_education,
            "experiences": _build_experience,
            "projects": _build_project,
        }[collection]
        return [builder(block) for block in blocks]
    builder_by_line = {
        "publications": _build_publication,
        "certificates": _build_certificate,
        "teaching": _build_teaching,
        "awards": _build_award,
        "languages": _build_language,
    }
    entries: list[dict[str, Any]] = []
    for line in lines:
        text = _BULLET_PREFIX.sub("", line).strip()
        if not text:
            continue
        if collection == "skills":
            entries.extend({"name": part} for part in _TOKEN_SPLIT.split(text) if part)
        else:
            entries.append(builder_by_line[collection](text))
    return entries


def _split_blocks(lines: list[str]) -> list[list[str]]:
    """空行或年份起始行视为新条目的开始。"""
    blocks: list[list[str]] = []
    block: list[str] = []
    for line in lines:
        text = _BULLET_PREFIX.sub("", line).strip()
        if not text:
            if block:
                blocks.append(block)
                block = []
            continue
        if block and _YEAR_LEAD.match(text):
            blocks.append(block)
            block = []
        block.append(text)
    if block:
        blocks.append(block)
    return blocks


# 教育/经历行的标签词：既是「标签-值对」的键，也不得被后缀启发式当成值。
# 「学校」本身以「学校」结尾（在 _SCHOOL_SUFFIXES 中），不排除会解析出 school=「学校」。
_EDU_LABEL_TOKENS = {"学校", "院校", "毕业院校", "专业", "学历", "学位"}
_EXP_ORG_LABEL_TOKENS = {"工作单位", "单位", "任职单位"}
_EXP_ROLE_LABEL_TOKENS = {"岗位", "职位", "职务"}
_LABEL_VALUE_PREFIX = re.compile(
    r"^(学校|院校|毕业院校|专业|学历|学位|工作单位|任职单位|单位|岗位|职位|职务)[：:](.+)$"
)


def _expand_label_value_tokens(tokens: list[str]) -> list[str]:
    """把「学校：北京大学」拆成「学校」「北京大学」两个 token，交给标签-值配对。"""
    expanded: list[str] = []
    for token in tokens:
        match = _LABEL_VALUE_PREFIX.match(token)
        expanded.extend([match.group(1), match.group(2)] if match else [token])
    return expanded


def _labeled_value(tokens: list[str], labels: set[str]) -> str:
    """标签（冒号或空白分隔）后紧跟的 token 是值：「学校 北京大学」→ 北京大学。"""
    for index, token in enumerate(tokens):
        if token.rstrip("：:") in labels and index + 1 < len(tokens):
            return tokens[index + 1]
    return ""


def _build_education(block: list[str]) -> dict[str, Any]:
    head, highlights = block[0], block[1:]
    start, end, remaining = _take_date_range(head)
    degree_match = _DEGREE.search(remaining)
    degree = degree_match.group(0) if degree_match else ""
    tokens = _expand_label_value_tokens(_tokens(remaining))
    school = _labeled_value(tokens, {"学校", "院校", "毕业院校"}) or _pick_token(
        [token for token in tokens if token.rstrip("：:") not in _EDU_LABEL_TOKENS], _SCHOOL_SUFFIXES
    )
    major = _labeled_value(tokens, {"专业"}) or next(
        (
            token
            for token in tokens
            if token not in {school, degree}
            and token.rstrip("：:") not in _EDU_LABEL_TOKENS
            and not _DEGREE.fullmatch(token)
            and 2 <= len(token) <= 20
        ),
        "",
    )
    return _entry(school=school, degree=degree, major=major, start=start, end=end, highlights=highlights)


def _build_experience(block: list[str]) -> dict[str, Any]:
    head, highlights = block[0], block[1:]
    start, end, remaining = _take_date_range(head)
    tokens = _expand_label_value_tokens(_tokens(remaining))
    organization = _labeled_value(tokens, _EXP_ORG_LABEL_TOKENS) or _pick_token(
        [token for token in tokens if token.rstrip("：:") not in _EXP_ORG_LABEL_TOKENS | _EXP_ROLE_LABEL_TOKENS],
        _ORG_SUFFIXES,
    )
    role = _labeled_value(tokens, _EXP_ROLE_LABEL_TOKENS) or _pick_token(
        [token for token in tokens if token != organization], _ROLE_SUFFIXES
    )
    return _entry(organization=organization, role=role, start=start, end=end, highlights=highlights)


def _build_project(block: list[str]) -> dict[str, Any]:
    head, highlights = block[0], block[1:]
    _, _, remaining = _take_date_range(head)
    tokens = _tokens(remaining)
    role = _pick_token(tokens, _PROJECT_ROLE_SUFFIXES)
    name = _pick_token([token for token in tokens if token != role], _PROJECT_NAME_HINTS)
    if not name:
        name = next((token for token in tokens if token != role), remaining)
    return _entry(name=name, role=role, highlights=highlights)


def _build_publication(line: str) -> dict[str, Any]:
    doi_match = _DOI.search(line)
    journal_match = _JOURNAL.search(line)
    year_match = _YEAR.search(line)
    return _entry(
        title=line,
        journal=journal_match.group(1) if journal_match else "",
        year=year_match.group(0) if year_match else "",
        doi=doi_match.group(0) if doi_match else "",
    )


def _build_certificate(line: str) -> dict[str, Any]:
    year_match = _YEAR.search(line)
    tokens = _tokens(line)
    issuer = _pick_token(tokens, _ISSUER_SUFFIXES)
    name = next((token for token in tokens if token != issuer and not _YEAR.fullmatch(token)), line)
    return _entry(name=name, issuer=issuer, year=year_match.group(0) if year_match else "")


def _build_teaching(line: str) -> dict[str, Any]:
    year_match = _YEAR.search(line)
    tokens = _tokens(line)
    institution = _pick_token(tokens, _SCHOOL_SUFFIXES + ["医院"])
    role = _pick_token([token for token in tokens if token != institution], ["主讲", "助教", "带教", "讲师", "授课"])
    course = next(
        (token for token in tokens if token not in {institution, role} and not _YEAR.fullmatch(token) and ("课" in token or "学" in token)),
        "",
    )
    return _entry(course=course, role=role, institution=institution, year=year_match.group(0) if year_match else "")


def _build_award(line: str) -> dict[str, Any]:
    year_match = _YEAR.search(line)
    level_match = _AWARD_LEVEL.search(line)
    tokens = _tokens(line)
    issuer = _pick_token(tokens, _ISSUER_SUFFIXES)
    name = next((token for token in tokens if token != issuer and not _YEAR.fullmatch(token)), line)
    return _entry(
        name=name,
        issuer=issuer,
        year=year_match.group(0) if year_match else "",
        level=level_match.group(0) if level_match else "",
    )


def _build_language(line: str) -> dict[str, Any]:
    name_match = _LANGUAGE_NAME.search(line)
    level_match = _LANGUAGE_LEVEL.search(line)
    return _entry(
        name=name_match.group(0) if name_match else line,
        level=normalize_text(level_match.group(0)) if level_match else "",
    )


def _take_date_range(text: str) -> tuple[str, str, str]:
    match = _DATE_RANGE.search(text)
    if not match:
        return "", "", text
    start = match.group("start").rstrip("月")
    end = match.group("end").rstrip("月")
    remaining = normalize_text(text[: match.start()] + " " + text[match.end() :])
    return start, end, remaining


def _tokens(text: str) -> list[str]:
    return [token for token in _TOKEN_SPLIT.split(text) if token]


def _pick_token(tokens: list[str], suffixes: list[str]) -> str:
    for token in tokens:
        if any(suffix in token for suffix in suffixes):
            return token
    return ""


def _entry(**fields: Any) -> dict[str, Any]:
    return {key: value for key, value in fields.items() if value}
