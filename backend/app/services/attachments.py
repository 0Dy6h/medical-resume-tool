from __future__ import annotations

from dataclasses import dataclass, replace
from io import BytesIO
from itertools import chain, islice
from typing import Any
from urllib.parse import urljoin, urlparse

from bs4 import BeautifulSoup
from openpyxl import load_workbook

from app.services.classifier import normalize_text
from app.services.crawler import ParsedJob, parse_job_from_text
from app.services.document_limits import DocumentLimitError, check_ooxml_archive


SUPPORTED_ATTACHMENT_EXTENSIONS = {".xlsx", ".xls", ".pdf"}
MAX_TABLE_ROWS = 5000
MAX_TABLE_COLUMNS = 128
MAX_TABLE_CELLS = 200_000


@dataclass(frozen=True)
class XlsxRow:
    sheet_name: str
    row_index: int
    headers: list[str]
    values: dict[str, str]
    raw_text: str


def extract_attachment_links(html: str, base_url: str) -> list[dict[str, str]]:
    soup = BeautifulSoup(html, "html.parser")
    attachments: list[dict[str, str]] = []
    seen: set[str] = set()
    for anchor in soup.find_all("a", href=True):
        href = str(anchor["href"]).strip()
        absolute_url = urljoin(base_url, href)
        extension = _extension_from_url(absolute_url)
        if extension not in SUPPORTED_ATTACHMENT_EXTENSIONS:
            continue
        if absolute_url in seen:
            continue
        seen.add(absolute_url)
        name = normalize_text(anchor.get_text(" ")) or absolute_url.rsplit("/", 1)[-1]
        attachments.append({"name": name, "url": absolute_url, "extension": extension})
    return attachments


def parse_xlsx_table(content: bytes) -> list[XlsxRow]:
    check_ooxml_archive(content)
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    rows: list[XlsxRow] = []
    visited_rows = 0
    visited_cells = 0
    try:
        for sheet in workbook.worksheets:
            if (sheet.max_row or 0) > MAX_TABLE_ROWS or (sheet.max_column or 0) > MAX_TABLE_COLUMNS:
                raise DocumentLimitError("附件表格行列过多，无法安全解析")
            iterator = sheet.iter_rows(max_row=MAX_TABLE_ROWS + 1, max_col=min(sheet.max_column or MAX_TABLE_COLUMNS, MAX_TABLE_COLUMNS), values_only=True)
            preview = list(islice(iterator, 10))
            if not preview:
                continue
            header_index, headers = _find_header_row(preview)
            if not headers:
                continue
            for offset, row in enumerate(chain(preview, iterator), start=1):
                # Ignore iterator padding after the advertised end of a sheet.
                if sheet.max_row is not None and offset > sheet.max_row:
                    break
                visited_rows += 1
                visited_cells += len(row)
                if visited_rows > MAX_TABLE_ROWS or visited_cells > MAX_TABLE_CELLS:
                    raise DocumentLimitError("附件表格内容过多，无法安全解析")
                if offset <= header_index + 1:
                    continue
                values = _row_values(headers, row)
                if not values:
                    continue
                raw_text = "；".join(f"{key}：{value}" for key, value in values.items())
                rows.append(
                    XlsxRow(sheet_name=str(sheet.title), row_index=offset, headers=headers,
                            values=values, raw_text=raw_text)
                )
    finally:
        workbook.close()
    return rows


def build_jobs_from_xlsx_attachment(
    *,
    content: bytes,
    attachment: dict[str, str],
    announcement_url: str,
    institution: dict[str, Any],
    parser_name: str,
    parser_warning: str | None = None,
) -> list[ParsedJob]:
    jobs: list[ParsedJob] = []
    for row in parse_xlsx_table(content):
        title = _pick(row.values, ["岗位名称", "招聘岗位", "岗位", "职位名称", "职位"])
        unit = _pick(row.values, ["招聘单位", "单位", "用人部门", "部门", "科室"])
        profession = _pick(row.values, ["专业", "所学专业", "专业要求"])
        education = _pick(row.values, ["学历", "学历要求", "学位", "学历学位"])
        if not title:
            title = _fallback_title(row.values, attachment["name"])
        body_parts = []
        if unit:
            body_parts.append(f"招聘单位：{unit}")
        if profession:
            body_parts.append(f"专业要求：{profession}")
        if education:
            body_parts.append(f"学历要求：{education}")
        body_parts.append(row.raw_text)
        job = parse_job_from_text(
            title=title,
            body="。".join(body_parts),
            source_url=f"{attachment['url']}#{row.sheet_name}-{row.row_index}",
            institution_type=institution["institution_type"],
            region=institution["region"],
            parser_name=parser_name,
            department=unit,
        )
        evidence = dict(job.extraction_evidence)
        evidence.update(
            {
                "announcement_url": announcement_url,
                "attachment_url": attachment["url"],
                "attachment_name": attachment["name"],
                "sheet_name": row.sheet_name,
                "row_index": row.row_index,
                "headers": row.headers,
            }
        )
        if parser_warning:
            evidence["parser_warning"] = parser_warning
        jobs.append(replace(job, extraction_evidence=evidence))
    return jobs


def status_for_attachment(
    attachment: dict[str, str],
    attachment_bytes_by_url: dict[str, bytes] | None,
    attachment_errors_by_url: dict[str, str] | None,
) -> dict[str, str]:
    """Map a discovered attachment to its fetch outcome (parsed / failed / discovered)."""
    if attachment["url"] in (attachment_bytes_by_url or {}):
        return attachment_status(attachment, "parsed")
    if attachment["url"] in (attachment_errors_by_url or {}):
        return attachment_status(attachment, "failed", (attachment_errors_by_url or {})[attachment["url"]])
    return attachment_status(attachment, "discovered")


def attachment_status(attachment: dict[str, str], status: str, error: str | None = None) -> dict[str, str]:
    payload = {
        "name": attachment["name"],
        "url": attachment["url"],
        "extension": attachment["extension"],
        "status": status,
    }
    if error:
        payload["error"] = error
    return payload


def _extension_from_url(url: str) -> str:
    path = urlparse(url).path.lower()
    for extension in SUPPORTED_ATTACHMENT_EXTENSIONS:
        if path.endswith(extension):
            return extension
    return ""


def _find_header_row(table: list[tuple[Any, ...]]) -> tuple[int, list[str]]:
    for index, row in enumerate(table[:10]):
        headers = [_cell_text(value) for value in row]
        # Keep physical column positions, including empty spacer/merged cells.
        # Removing blank headings shifts every following fact to the wrong key.
        if sum(bool(value) for value in headers) >= 2 and any(_is_job_header(value) for value in headers):
            return index, headers
    first = [_cell_text(value) for value in table[0]]
    return (0, first) if sum(bool(value) for value in first) >= 2 else (0, [])


def _row_values(headers: list[str], row: tuple[Any, ...]) -> dict[str, str]:
    values: dict[str, str] = {}
    for header, value in zip(headers, row):
        if not header:
            continue
        text = _cell_text(value)
        if text:
            values[header] = text
    return values


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return normalize_text(str(value))


def _is_job_header(value: str) -> bool:
    return any(keyword in value for keyword in ["岗位", "职位", "专业", "学历", "招聘单位", "科室"])


def _pick(values: dict[str, str], candidates: list[str]) -> str:
    for candidate in candidates:
        for key, value in values.items():
            if candidate in key and value:
                return value
    return ""


def _fallback_title(values: dict[str, str], attachment_name: str) -> str:
    for key, value in values.items():
        if any(keyword in key for keyword in ["岗位", "职位", "专业"]):
            return value
    first_value = next(iter(values.values()), "")
    return first_value or attachment_name.replace(".xlsx", "").replace(".xls", "")
